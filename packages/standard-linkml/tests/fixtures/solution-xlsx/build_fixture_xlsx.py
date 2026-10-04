"""Generate fixture workbook with deliberate defects for solution-xlsx tests."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def build(path: Path) -> Path:
    wb = Workbook()
    # --- Object ---
    ws = wb.active
    ws.title = "Object"
    ws.append([None] * 6)
    ws.append(
        [
            "SrcSystem",
            "ObjectCode",
            "ObjectName",
            "SrcObjectCode",
            "SrcObjectName",
            "SrcDescription",
        ]
    )
    ws.append(
        ["MDM", "ENTERPRISE", "ЮЛ", "V_ENTERPRISE", "ЮЛ", "Legal entity"]
    )
    ws.append(["MDM", "PERSON", "ФЛ", "V_PERSON", "ФЛ", "Person"])
    # CRM Object uses UPPER case
    ws.append(["CRM", "CUSTOMER", "Клиент", "ACCOUNT", "Клиент", None])
    ws.append(["CRM", "CONTACT", "Контакт", "CONTACT", "Контакт", None])
    # ЕСЭД blank row
    ws.append(["ЕСЭД", None, None, None, None, None])

    # --- ObjectAttribute (header on row 6, 1-based) ---
    wa = wb.create_sheet("ObjectAttribute")
    wa.append([None, "Комментарии"])
    wa.append([])
    wa.append([])
    wa.append([None, "Target model attributes"])
    wa.append([])
    wa.append(
        [
            "SortOrder",
            "ObjectCode",
            "ObjectName",
            "AttributeCode",
            "AttributeDescription",
            "DataType",
            "Mandatory",
            "PK",
            "AK",
            "FK",
            "SrcSystem",
            "SrcObjectCode",
            "SrcObjectName",
            "SrcAttributeCode",
            "SrcAttributeDescription",
            "SrcDataType",
            "Comments",
            "BaseSrcSCode",
            "BaseSrcSName",
            "BaseSrcObjectCode",
        ]
    )
    # MDM ENTERPRISE
    wa.append(
        [
            0,
            "ENTERPRISE",
            "ЮЛ",
            "ENTERPRISE_ID",
            "ID",
            "VARCHAR2",
            "Y",
            "Y",
            None,
            None,
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "GOLD_ID",
            "ID",
            "VARCHAR2",
            None,
            "MDM",
            "MDM",
            "V_ENTERPRISE",
        ]
    )
    wa.append(
        [
            1,
            "ENTERPRISE",
            "ЮЛ",
            "INN",
            "ИНН",
            "VARCHAR2",
            None,
            None,
            None,
            None,
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "INN",
            "ИНН",
            "VARCHAR2",
            None,
            "FTS",
            "ФНС",
            "V_FTS",
        ]
    )
    # MDM PERSON + nbsp attr
    wa.append(
        [
            0,
            "PERSON",
            "ФЛ",
            "PERSON_ID",
            "ID",
            "VARCHAR2",
            "Y",
            "Y",
            None,
            None,
            "MDM",
            "V_PERSON",
            "ФЛ",
            "GOLD_ID",
            "ID",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    wa.append(
        [
            1,
            "PERSON",
            "ФЛ",
            "LAST_NAME\xa0",
            "Фамилия",
            "VARCHAR2",
            None,
            None,
            None,
            None,
            "MDM",
            "V_PERSON",
            "ФЛ",
            "LAST_NAME",
            "Фамилия",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    # MDM link + FK
    wa.append(
        [
            0,
            "PERSON_CEO",
            "CEO link",
            "ENTERPRISE_ID",
            "FK ent",
            "VARCHAR2",
            "Y",
            "Y",
            None,
            "ENTERPRISE.ENTERPRISE_ID",
            "MDM",
            "T_CEO",
            "CEO",
            "ENTERPRISE_ID",
            "FK",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    wa.append(
        [
            1,
            "PERSON_CEO",
            "CEO link",
            "PERSON_ID",
            "FK person",
            "VARCHAR2",
            "Y",
            "Y",
            None,
            "PERSON.PERSON_ID",
            "MDM",
            "T_CEO",
            "CEO",
            "PERSON_ID",
            "FK",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    # src-only row (no target)
    wa.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "ADDRESS_F",
            "addr",
            "VARCHAR2",
            None,
            "ECBD",
            "ЕГРЮЛ",
            "V_ECBD",
        ]
    )
    # duplicate attr
    wa.append(
        [
            2,
            "ENTERPRISE",
            "ЮЛ",
            "INN",
            "ИНН dup",
            "VARCHAR2",
            None,
            None,
            None,
            None,
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "INN2",
            "ИНН2",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    # unknown type
    wa.append(
        [
            3,
            "ENTERPRISE",
            "ЮЛ",
            "WEIRD",
            "weird",
            "FOOBAR99",
            None,
            None,
            None,
            None,
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "WEIRD",
            "weird",
            "FOOBAR99",
            None,
            None,
            None,
            None,
        ]
    )
    # broken FK
    wa.append(
        [
            4,
            "ENTERPRISE",
            "ЮЛ",
            "BAD_FK",
            "bad",
            "VARCHAR2",
            None,
            None,
            None,
            "MISSING.ID",
            "MDM",
            "V_ENTERPRISE",
            "ЮЛ",
            "BAD_FK",
            "bad",
            "VARCHAR2",
            None,
            None,
            None,
            None,
        ]
    )
    # CRM case mismatch Customer vs CUSTOMER
    wa.append(
        [
            0,
            "Customer",
            "Клиент",
            "Id",
            "ID",
            "uuid",
            "Y",
            "Y",
            None,
            None,
            "CRM",
            "ACCOUNT",
            "Клиент",
            "accountid",
            "ID",
            "uniqueidentifier",
            None,
            None,
            None,
            None,
        ]
    )
    wa.append(
        [
            1,
            "Customer",
            "Клиент",
            "PrimaryContact",
            "Contact",
            "uuid",
            None,
            None,
            None,
            "Contact.Id",
            "CRM",
            "ACCOUNT",
            "Клиент",
            "primarycontactid",
            "Contact",
            "uniqueidentifier",
            None,
            None,
            None,
            None,
        ]
    )
    wa.append(
        [
            0,
            "Contact",
            "Контакт",
            "Id",
            "ID",
            "uuid",
            "Y",
            "Y",
            None,
            None,
            "CRM",
            "CONTACT",
            "Контакт",
            "contactid",
            "ID",
            "uniqueidentifier",
            None,
            None,
            None,
            None,
        ]
    )
    # ЕСЭД synthesized objects
    wa.append(
        [
            0,
            "document",
            "Документ",
            "document_id",
            "ID",
            "UUID",
            "Y",
            "Y",
            None,
            "instance.instance_id",
            "ЕСЭД",
            "document",
            "Документ",
            "document_id",
            "ID",
            "UUID",
            None,
            None,
            None,
            None,
        ]
    )
    wa.append(
        [
            0,
            "instance",
            "Экземпляр",
            "instance_id",
            "ID",
            "UUID",
            "Y",
            "Y",
            None,
            None,
            "ЕСЭД",
            "instance",
            "Экземпляр",
            "instance_id",
            "ID",
            "UUID",
            None,
            None,
            None,
            None,
        ]
    )

    # ignored sheets
    wb.create_sheet("1")
    wb.create_sheet("Блок api")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    root = Path(__file__).resolve().parent
    build(root / "fixture.xlsx")
    print("wrote", root / "fixture.xlsx")
