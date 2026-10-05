"""Load ER-dictionary sheets from .xlsx or a directory of CSV files."""

from __future__ import annotations

import csv
import logging
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from moex_standard_linkml.ingest.profile import IngestProfile, SheetSpec

logger = logging.getLogger(__name__)


@dataclass
class SheetTable:
    """One logical sheet: rows as dicts keyed by logical column names."""

    logical_name: str
    source_name: str
    rows: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class WorkbookTables:
    entities: SheetTable
    attributes: SheetTable
    relationships: SheetTable | None
    conceptual: SheetTable | None = None
    data_carriers: SheetTable | None = None
    physical_fields: SheetTable | None = None
    mappings: SheetTable | None = None
    warnings: list[str] = field(default_factory=list)


class WorkbookError(ValueError):
    """Missing required sheet or unreadable workbook."""


def load_workbook_tables(
    source: Path | str,
    profile: IngestProfile,
) -> WorkbookTables:
    source = Path(source)
    if source.is_dir():
        raw = _load_csv_dir(source)
    elif source.suffix.lower() in {".xlsx", ".xlsm"}:
        raw = _load_xlsx(source)
    else:
        raise WorkbookError(
            f"Unsupported workbook source (need .xlsx or CSV directory): {source}"
        )

    warnings_all: list[str] = []
    entities = _project_sheet("entities", profile.sheets.entities, raw, warnings_all)
    attributes = _project_sheet(
        "attributes", profile.sheets.attributes, raw, warnings_all
    )
    relationships = _load_optional_sheet(
        "relationships", profile.sheets.relationships, raw, warnings_all
    )
    conceptual = _load_optional_sheet(
        "conceptual", profile.sheets.conceptual, raw, warnings_all
    )
    data_carriers = _load_optional_sheet(
        "data_carriers", profile.sheets.data_carriers, raw, warnings_all
    )
    physical_fields = _load_optional_sheet(
        "physical_fields", profile.sheets.physical_fields, raw, warnings_all
    )
    mappings = _load_optional_sheet(
        "mappings", profile.sheets.mappings, raw, warnings_all
    )

    return WorkbookTables(
        entities=entities,
        attributes=attributes,
        relationships=relationships,
        conceptual=conceptual,
        data_carriers=data_carriers,
        physical_fields=physical_fields,
        mappings=mappings,
        warnings=warnings_all,
    )


def _load_optional_sheet(
    logical_name: str,
    spec: SheetSpec | None,
    raw: dict[str, list[dict[str, Any]]],
    warnings_all: list[str],
) -> SheetTable | None:
    if spec is None:
        return None
    if _sheet_present(spec.sheet, raw) or spec.required:
        return _project_sheet(logical_name, spec, raw, warnings_all)
    warnings_all.append(
        f"Optional sheet '{spec.sheet}' not found; {logical_name} empty"
    )
    return SheetTable(
        logical_name=logical_name,
        source_name=spec.sheet,
        rows=[],
    )


def _sheet_present(sheet: str, raw: dict[str, list[dict[str, Any]]]) -> bool:
    if sheet in raw:
        return True
    return any(k.lower() == sheet.lower() for k in raw)


def _load_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                result[sheet_name] = []
                continue
            headers = [_cell_str(h) for h in header_row]
            if not any(headers):
                result[sheet_name] = []
                continue
            rows: list[dict[str, Any]] = []
            for raw_row in rows_iter:
                if raw_row is None or all(
                    c is None or str(c).strip() == "" for c in raw_row
                ):
                    continue
                row: dict[str, Any] = {}
                for i, key in enumerate(headers):
                    if not key:
                        continue
                    val = raw_row[i] if i < len(raw_row) else None
                    row[key] = _normalize_cell(val)
                rows.append(row)
            result[sheet_name] = rows
    finally:
        wb.close()
    return result


def _load_csv_dir(directory: Path) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() != ".csv" or not path.is_file():
            continue
        with path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            if reader.fieldnames is None:
                result[path.stem] = []
                continue
            rows: list[dict[str, Any]] = []
            for raw in reader:
                if all(
                    v is None or str(v).strip() == "" for v in raw.values()
                ):
                    continue
                rows.append(
                    {
                        (k or "").strip(): _normalize_cell(v)
                        for k, v in raw.items()
                        if k is not None and str(k).strip()
                    }
                )
            result[path.stem] = rows
    return result


def _project_sheet(
    logical_name: str,
    spec: SheetSpec,
    raw: dict[str, list[dict[str, Any]]],
    warnings_all: list[str],
) -> SheetTable:
    if spec.sheet not in raw:
        match = next(
            (k for k in raw if k.lower() == spec.sheet.lower()),
            None,
        )
        if match is None:
            if spec.required:
                available = ", ".join(sorted(raw)) or "(none)"
                raise WorkbookError(
                    f"Required sheet '{spec.sheet}' not found "
                    f"(logical={logical_name}). Available: {available}"
                )
            return SheetTable(
                logical_name=logical_name,
                source_name=spec.sheet,
                rows=[],
            )
        source_name = match
    else:
        source_name = spec.sheet

    source_rows = raw[source_name]
    col_map = _logical_column_map(spec)
    expected_headers = set(col_map.values())
    unknown: set[str] = set()
    projected: list[dict[str, Any]] = []

    for row in source_rows:
        for header in row:
            if header and header not in expected_headers:
                unknown.add(header)
        projected.append(
            {logical: row.get(header) for logical, header in col_map.items()}
        )

    sheet_warnings: list[str] = []
    for header in sorted(unknown):
        msg = (
            f"Sheet '{source_name}': unknown column '{header}' ignored "
            f"(logical={logical_name})"
        )
        sheet_warnings.append(msg)
        warnings_all.append(msg)
        warnings.warn(msg, stacklevel=2)
        logger.warning(msg)

    if not projected and spec.required:
        raise WorkbookError(
            f"Required sheet '{source_name}' has no data rows (logical={logical_name})"
        )

    return SheetTable(
        logical_name=logical_name,
        source_name=source_name,
        rows=projected,
        warnings=sheet_warnings,
    )


def _logical_column_map(spec: SheetSpec) -> dict[str, str]:
    """Map logical field name → workbook header (skip unset optional columns)."""
    data = spec.columns.model_dump(exclude_none=True)
    return {k: str(v) for k, v in data.items() if v}


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    return text if text else None


def write_csv_dir_as_xlsx(
    csv_dir: Path,
    xlsx_path: Path,
    *,
    sheet_names: dict[str, str] | None = None,
) -> Path:
    """Helper for tests: CSV stem → workbook sheet (optional rename)."""
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    sheet_names = sheet_names or {}
    for path in sorted(csv_dir.glob("*.csv")):
        title = sheet_names.get(path.stem, path.stem)
        ws = wb.create_sheet(title=title)
        with path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            for row in reader:
                ws.append(row)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return xlsx_path
