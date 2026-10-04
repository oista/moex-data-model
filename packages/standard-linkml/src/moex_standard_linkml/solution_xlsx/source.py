"""Read Object / ObjectAttribute sheets into SolutionIR for one SrcSystem."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from moex_standard_linkml.solution_xlsx.diagnostics import Diagnostic, Severity, SourceRef
from moex_standard_linkml.solution_xlsx.ir import (
    AttributeDef,
    ObjectDef,
    SolutionIR,
    SrcOnlyRow,
)
from moex_standard_linkml.solution_xlsx.normalize import match_key, normalize_code
from moex_standard_linkml.solution_xlsx.profile import SolutionXlsxProfile


class SourceError(ValueError):
    """Unreadable workbook or missing required sheet/columns."""


def _has_cyrillic(text: str) -> bool:
    return any("\u0400" <= c <= "\u04FF" for c in text)


def _looks_technical_code(text: str) -> bool:
    """ASCII identifier-like token (Id, MoexINN, account_id)."""
    if not text or _has_cyrillic(text):
        return False
    cleaned = text.replace("_", "").replace("-", "")
    return cleaned.isalnum() and any(c.isalpha() for c in cleaned)


def _pick_attribute_names(
    code: str, description: str | None
) -> tuple[str, str | None, bool]:
    """
    Return (technical_name, title, swapped).

    Some CRM rows put RU label in AttributeCode and EN Id in AttributeDescription.
    """
    if (
        description
        and _has_cyrillic(code)
        and _looks_technical_code(description)
    ):
        return description, code, True
    return code, description, False


def _as_bool(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().lower()
    return text in {"1", "true", "yes", "y", "x", "pk"}


def _cell(row: tuple[Any, ...], idx: int | None) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _find_header_row(
    rows: list[tuple[Any, ...]],
    required_headers: list[str],
    *,
    max_scan: int = 30,
) -> tuple[int, dict[str, int]]:
    """Return 0-based index of header row and map header→col index."""
    needed = {h.strip() for h in required_headers if h and h.strip()}
    scan = rows[:max_scan] if max_scan else rows
    for i, row in enumerate(scan):
        headers: dict[str, int] = {}
        for j, cell in enumerate(row):
            if cell is None:
                continue
            name = str(cell).strip()
            if name:
                headers[name] = j
        if needed and needed.issubset(headers.keys()):
            return i, headers
        if not needed and headers:
            return i, headers
    available = ""
    if rows:
        sample = [str(c).strip() for c in rows[0] if c is not None][:12]
        available = ", ".join(sample)
    raise SourceError(
        f"Header row with columns {sorted(needed)} not found "
        f"(scanned {len(scan)} rows). First-row sample: {available or '(empty)'}"
    )


def _load_sheet_rows(path: Path, sheet_name: str) -> list[tuple[Any, ...]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = {n.casefold(): n for n in wb.sheetnames}
        key = sheet_name.casefold()
        if key not in names:
            raise SourceError(
                f"Sheet '{sheet_name}' not found. Available: {', '.join(wb.sheetnames)}"
            )
        ws = wb[names[key]]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_solution_ir(
    xlsx_path: Path | str,
    profile: SolutionXlsxProfile,
    *,
    src_system: str,
) -> SolutionIR:
    """Load and filter workbook for one system into SolutionIR (+ early diagnostics)."""
    path = Path(xlsx_path)
    if not path.is_file():
        raise SourceError(f"xlsx not found: {path}")

    sys_cfg = profile.system_for(src_system)
    if sys_cfg is None:
        raise SourceError(
            f"SrcSystem '{src_system}' is not defined in profile.systems"
        )

    cols = profile.columns
    obj_spec = profile.sheets.objects
    attr_spec = profile.sheets.object_attributes

    diagnostics: list[Diagnostic] = []
    ignored: list[str] = []

    # Detect ignored sheets present in workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        present = set(wb.sheetnames)
    finally:
        wb.close()
    for name in profile.ignored_sheets:
        if name in present or any(n.casefold() == name.casefold() for n in present):
            ignored.append(name)
    if ignored:
        diagnostics.append(
            Diagnostic(
                code="SXI-SRC-011",
                severity=Severity.INFO,
                message_ru=(
                    "Листы вне объёма v1 присутствуют в книге и пропущены: "
                    + ", ".join(ignored)
                ),
                remediation="Импорт использует только Object и ObjectAttribute.",
                source_ref=SourceRef(sheet="(workbook)", row=0),
            )
        )

    try:
        object_rows = _load_sheet_rows(path, obj_spec.sheet)
        attr_rows = _load_sheet_rows(path, attr_spec.sheet)
    except SourceError as exc:
        diagnostics.append(
            Diagnostic(
                code="SXI-SRC-001",
                severity=Severity.ERROR,
                message_ru=str(exc),
                remediation="Проверьте путь к xlsx и имена листов в профиле.",
                source_ref=SourceRef(sheet=obj_spec.sheet, row=0),
            )
        )
        return SolutionIR(
            src_system=sys_cfg.src_system,
            diagnostics=diagnostics,
            ignored_sheets=ignored,
            source_filename=path.name,
            source_sha256=_sha256_file(path),
        )

    try:
        obj_header_i, obj_headers = _find_header_row(
            object_rows, obj_spec.required_headers
        )
        attr_header_i, attr_headers = _find_header_row(
            attr_rows, attr_spec.required_headers
        )
    except SourceError as exc:
        diagnostics.append(
            Diagnostic(
                code="SXI-SRC-001",
                severity=Severity.ERROR,
                message_ru=str(exc),
                remediation=(
                    "Убедитесь, что строка заголовка содержит ожидаемые колонки "
                    "(см. profile.columns / required_headers)."
                ),
                source_ref=SourceRef(sheet=attr_spec.sheet, row=1),
            )
        )
        return SolutionIR(
            src_system=sys_cfg.src_system,
            diagnostics=diagnostics,
            ignored_sheets=ignored,
            source_filename=path.name,
            source_sha256=_sha256_file(path),
        )

    # Missing optional columns → None indices (not fatal)
    def idx(headers: dict[str, int], header_name: str) -> int | None:
        return headers.get(header_name)

    ir = SolutionIR(
        src_system=sys_cfg.src_system,
        diagnostics=diagnostics,
        ignored_sheets=ignored,
        source_filename=path.name,
        source_sha256=_sha256_file(path),
    )

    system_key = sys_cfg.src_system.casefold()

    # --- Object sheet ---
    for offset, row in enumerate(object_rows[obj_header_i + 1 :], start=obj_header_i + 2):
        sys_raw = _cell(row, idx(obj_headers, cols.src_system))
        sys_n = normalize_code(sys_raw)
        if sys_n.value is None:
            # empty system — may be ЕСЭД blank row
            code_n = normalize_code(_cell(row, idx(obj_headers, cols.object_code)))
            if code_n.value is None:
                # Fully empty object row for this scan — if any cell has system-like empty
                # but row has content in other cols belonging to our filter later
                if all(
                    (_cell(row, j) is None or str(_cell(row, j)).strip() == "")
                    for j in range(len(row))
                ):
                    continue
                # Blank Object row (ЕСЭД pattern)
                if match_key(sys_cfg.src_system) and any(
                    _cell(row, j) is not None for j in range(min(3, len(row)))
                ):
                    # Only emit once-style warning when SrcSystem empty and codes empty
                    ir.diagnostics.append(
                        Diagnostic(
                            code="SXI-OBJ-002",
                            severity=Severity.WARNING,
                            message_ru=(
                                f"Пустая строка на листе Object для контекста "
                                f"{sys_cfg.src_system}; объекты будут синтезированы "
                                f"из ObjectAttribute."
                            ),
                            remediation=(
                                "Заполните ObjectCode/ObjectName на листе Object "
                                "или оставьте синтез из ObjectAttribute."
                            ),
                            source_ref=SourceRef(
                                sheet=obj_spec.sheet,
                                row=offset,
                                column=cols.object_code,
                            ),
                        )
                    )
                continue
            continue

        if sys_n.value.casefold() != system_key:
            continue

        code_n = normalize_code(_cell(row, idx(obj_headers, cols.object_code)))
        if code_n.value is None:
            ir.diagnostics.append(
                Diagnostic(
                    code="SXI-OBJ-002",
                    severity=Severity.WARNING,
                    message_ru=(
                        f"Строка Object с SrcSystem={sys_n.value} без ObjectCode; "
                        f"пропущена."
                    ),
                    remediation="Укажите ObjectCode или удалите строку.",
                    source_ref=SourceRef(
                        sheet=obj_spec.sheet, row=offset, column=cols.object_code
                    ),
                )
            )
            continue

        name_n = normalize_code(_cell(row, idx(obj_headers, cols.object_name)))
        src_code_n = normalize_code(_cell(row, idx(obj_headers, cols.src_object_code)))
        src_name_n = normalize_code(_cell(row, idx(obj_headers, cols.src_object_name)))
        src_desc = _cell(row, idx(obj_headers, cols.src_description))
        src_desc_s = str(src_desc).strip() if src_desc is not None else None

        key = match_key(code_n.value)
        assert key is not None
        ir.objects[key] = ObjectDef(
            src_system=sys_cfg.src_system,
            object_code=code_n.value,
            object_name=name_n.value,
            src_object_code=src_code_n.value,
            src_object_name=src_name_n.value,
            src_description=src_desc_s or None,
            source_ref=SourceRef(
                sheet=obj_spec.sheet, row=offset, column=cols.object_code
            ),
        )

    # --- ObjectAttribute sheet ---
    for offset, row in enumerate(attr_rows[attr_header_i + 1 :], start=attr_header_i + 2):
        sys_n = normalize_code(_cell(row, idx(attr_headers, cols.src_system)))
        if sys_n.value is None:
            # skip completely empty
            if all(
                _cell(row, j) is None or str(_cell(row, j)).strip() == ""
                for j in range(len(row))
            ):
                continue
            ir.diagnostics.append(
                Diagnostic(
                    code="SXI-SRC-002",
                    severity=Severity.ERROR,
                    message_ru="Пустой SrcSystem в строке ObjectAttribute.",
                    remediation="Заполните SrcSystem (MDM/UCD/CRM/ЕСЭД).",
                    source_ref=SourceRef(
                        sheet=attr_spec.sheet, row=offset, column=cols.src_system
                    ),
                )
            )
            continue

        if sys_n.value.casefold() != system_key:
            # Unknown system globally?
            if sys_n.value.casefold() not in profile.known_systems():
                # Only report when scanning — but we filter per system; skip
                pass
            continue

        obj_n = normalize_code(_cell(row, idx(attr_headers, cols.object_code)))
        attr_n = normalize_code(_cell(row, idx(attr_headers, cols.attribute_code)))
        has_target = obj_n.value is not None and attr_n.value is not None

        if not has_target:
            ir.src_only.append(
                SrcOnlyRow(
                    src_system=sys_cfg.src_system,
                    src_object_code=normalize_code(
                        _cell(row, idx(attr_headers, cols.src_object_code))
                    ).value,
                    src_object_name=normalize_code(
                        _cell(row, idx(attr_headers, cols.src_object_name))
                    ).value,
                    src_attribute_code=normalize_code(
                        _cell(row, idx(attr_headers, cols.src_attribute_code))
                    ).value,
                    src_attribute_description=normalize_code(
                        _cell(row, idx(attr_headers, cols.src_attribute_description))
                    ).value,
                    src_data_type=normalize_code(
                        _cell(row, idx(attr_headers, cols.src_data_type))
                    ).value,
                    base_src_scode=normalize_code(
                        _cell(row, idx(attr_headers, cols.base_src_scode))
                    ).value,
                    base_src_sname=normalize_code(
                        _cell(row, idx(attr_headers, cols.base_src_sname))
                    ).value,
                    base_src_object_code=normalize_code(
                        _cell(row, idx(attr_headers, cols.base_src_object_code))
                    ).value,
                    source_ref=SourceRef(
                        sheet=attr_spec.sheet,
                        row=offset,
                        column=cols.src_attribute_code,
                    ),
                )
            )
            continue

        # Target row — resolve canonical object code from Object sheet
        obj_key = match_key(obj_n.value)
        assert obj_key is not None
        case_matched = False
        canonical_obj = obj_n.value
        if obj_key in ir.objects:
            canonical_obj = ir.objects[obj_key].object_code
            if canonical_obj != obj_n.value:
                case_matched = True
                ir.objects[obj_key].case_matched = True
        else:
            # Synthesize object later in rules; keep spelling from Attribute
            pass

        code_normalized = attr_n.changed or attr_n.had_nbsp or attr_n.had_whitespace
        if obj_n.changed or obj_n.had_nbsp or obj_n.had_whitespace:
            code_normalized = True

        sort_raw = _cell(row, idx(attr_headers, cols.sort_order))
        sort_order: int | None
        try:
            sort_order = int(sort_raw) if sort_raw is not None and str(sort_raw).strip() else None
        except (TypeError, ValueError):
            sort_order = None

        fk_n = normalize_code(_cell(row, idx(attr_headers, cols.fk)))
        dt_n = normalize_code(_cell(row, idx(attr_headers, cols.data_type)))
        src_attr_n = normalize_code(
            _cell(row, idx(attr_headers, cols.src_attribute_code))
        )
        src_obj_n = normalize_code(_cell(row, idx(attr_headers, cols.src_object_code)))

        desc_raw = _cell(row, idx(attr_headers, cols.attribute_description))
        desc = str(desc_raw).strip() if desc_raw is not None and str(desc_raw).strip() else None
        tech_name, title, swapped = _pick_attribute_names(attr_n.value, desc)  # type: ignore[arg-type]
        if swapped:
            code_normalized = True
            ir.diagnostics.append(
                Diagnostic(
                    code="SXI-ATTR-001",
                    severity=Severity.WARNING,
                    message_ru=(
                        f"AttributeCode '{attr_n.value}' похож на подпись, "
                        f"техническое имя взято из AttributeDescription "
                        f"'{tech_name}'."
                    ),
                    remediation=(
                        "Поменяйте местами AttributeCode и AttributeDescription "
                        "в xlsx (технический код — в AttributeCode)."
                    ),
                    source_ref=SourceRef(
                        sheet=attr_spec.sheet,
                        row=offset,
                        column=cols.attribute_code,
                    ),
                )
            )

        ir.attributes.append(
            AttributeDef(
                src_system=sys_cfg.src_system,
                object_code=canonical_obj,
                attribute_code=tech_name,
                attribute_description=title or desc,
                data_type=dt_n.value,
                mandatory=_as_bool(_cell(row, idx(attr_headers, cols.mandatory))),
                pk=_as_bool(_cell(row, idx(attr_headers, cols.pk))),
                ak=_as_bool(_cell(row, idx(attr_headers, cols.ak))),
                fk=fk_n.value,
                src_object_code=src_obj_n.value,
                src_attribute_code=src_attr_n.value,
                src_attribute_description=normalize_code(
                    _cell(row, idx(attr_headers, cols.src_attribute_description))
                ).value,
                src_data_type=normalize_code(
                    _cell(row, idx(attr_headers, cols.src_data_type))
                ).value,
                comments=normalize_code(
                    _cell(row, idx(attr_headers, cols.comments))
                ).value,
                base_src_scode=normalize_code(
                    _cell(row, idx(attr_headers, cols.base_src_scode))
                ).value,
                base_src_sname=normalize_code(
                    _cell(row, idx(attr_headers, cols.base_src_sname))
                ).value,
                base_src_object_code=normalize_code(
                    _cell(row, idx(attr_headers, cols.base_src_object_code))
                ).value,
                source_ref=SourceRef(
                    sheet=attr_spec.sheet,
                    row=offset,
                    column=cols.attribute_code,
                ),
                sort_order=sort_order,
                code_normalized=code_normalized,
            )
        )

        if case_matched:
            ir.diagnostics.append(
                Diagnostic(
                    code="SXI-OBJ-003",
                    severity=Severity.WARNING,
                    message_ru=(
                        f"ObjectCode '{obj_n.value}' совпал с Object только "
                        f"без учёта регистра → канон '{canonical_obj}'."
                    ),
                    remediation="Выровняйте регистр ObjectCode между листами.",
                    source_ref=SourceRef(
                        sheet=attr_spec.sheet, row=offset, column=cols.object_code
                    ),
                )
            )

    return ir
