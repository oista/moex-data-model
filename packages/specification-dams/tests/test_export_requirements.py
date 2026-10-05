"""RequirementCatalog → XLSX export (ADR-013)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from moex_dams.application.export_requirements import (
    EXPORT_COLUMNS,
    HEADER_LABELS,
    collect_requirement_rows,
    iter_catalog_rows,
    requirement_row,
    write_requirements_xlsx,
)

FIXTURE = Path(__file__).parent / "fixtures" / "requirements" / "bad-missing-statement.yaml"
IT_SOLUTION = (
    Path(__file__).resolve().parents[3]
    / "model-assets/specifications/moex-dams/0.1/requirements/it-solution-requirements.yaml"
)
CONCEPTUAL = (
    Path(__file__).resolve().parents[3]
    / "model-assets/specifications/moex-dams/0.1/requirements/conceptual-model-requirements.yaml"
)


def test_requirement_row_flattens_checks() -> None:
    row = requirement_row(
        {
            "element_id": "dams:req/GEN-001",
            "code": "GEN-001",
            "name": "model_package_identity",
            "title": "Идентичность",
            "description": "short",
            "statement": "У пакета должен быть идентификатор.",
            "requirement_level": "it_solution",
            "requirement_section": "GEN",
            "lifecycle_status": "approved",
            "applies_to": {
                "applies_target_class": "ModelPackage",
                "applies_target_kinds": ["table", "topic"],
                "applies_implementation_scope": "solution",
            },
            "formal_checks": [
                {
                    "check_id": "GEN-001.c1",
                    "kind": "slot_required",
                    "severity": "error",
                    "diagnostic_code": "DAMS-STRUCT-001",
                },
                {
                    "check_id": "GEN-001.c2",
                    "kind": "slot_required",
                    "severity": "warning",
                },
            ],
        },
        catalog_id="dams:req-catalog/it-solution/0.1",
        catalog_name="it_solution_requirements",
        source_path="it-solution-requirements.yaml",
    )
    assert row["code"] == "GEN-001"
    assert row["statement"] == "У пакета должен быть идентификатор."
    assert row["applies_target_kinds"] == "table; topic"
    assert row["check_count"] == "2"
    assert row["check_ids"] == "GEN-001.c1; GEN-001.c2"
    assert row["severities"] == "error; warning"
    assert row["diagnostic_codes"] == "DAMS-STRUCT-001; "
    assert list(row) == list(EXPORT_COLUMNS)


def test_iter_catalog_rows_reads_fixture() -> None:
    rows = iter_catalog_rows(FIXTURE, source_path="bad-missing-statement.yaml")
    assert len(rows) == 1
    assert rows[0]["code"] == "GEN-901"
    assert rows[0]["catalog_id"] == "dams:req-catalog/bad-missing-statement/0.1"


def test_write_xlsx_roundtrip(tmp_path: Path) -> None:
    rows = iter_catalog_rows(FIXTURE, source_path="fix.yaml")
    out = tmp_path / "req.xlsx"
    n = write_requirements_xlsx(rows, out)
    assert n == 1
    assert out.is_file()
    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "requirements"
    assert ws["A1"].value == HEADER_LABELS["catalog_id"]
    assert ws["E2"].value == "GEN-901"
    assert (ws["I2"].value or "").strip() == ""
    assert ws.freeze_panes == "A2"


def test_real_catalogs_row_count() -> None:
    root = Path(__file__).resolve().parents[3]
    rows = collect_requirement_rows([IT_SOLUTION, CONCEPTUAL], root=root)
    codes = [r["code"] for r in rows]
    assert "GEN-001" in codes
    assert "CM-GEN-001" in codes
    assert len(rows) == 53
    assert rows == sorted(
        rows,
        key=lambda r: (
            r["requirement_level"],
            r["requirement_section"],
            r["code"],
            r["element_id"],
        ),
    )
