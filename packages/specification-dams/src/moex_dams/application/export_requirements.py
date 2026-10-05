"""One-way RequirementCatalog YAML → XLSX projection (ADR-013)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

EXPORT_COLUMNS: tuple[str, ...] = (
    "catalog_id",
    "catalog_name",
    "source_path",
    "element_id",
    "code",
    "name",
    "title",
    "description",
    "statement",
    "requirement_level",
    "requirement_section",
    "lifecycle_status",
    "applies_target_class",
    "applies_target_kinds",
    "applies_implementation_scope",
    "applies_dams_model_level",
    "applies_implementation_profile",
    "check_count",
    "check_ids",
    "check_kinds",
    "severities",
    "diagnostic_codes",
)

# Human-readable header labels for Excel review.
HEADER_LABELS: dict[str, str] = {
    "catalog_id": "Каталог ID",
    "catalog_name": "Каталог",
    "source_path": "Источник",
    "element_id": "Element ID",
    "code": "Код",
    "name": "Имя",
    "title": "Заголовок",
    "description": "Описание",
    "statement": "Формулировка",
    "requirement_level": "Уровень",
    "requirement_section": "Раздел",
    "lifecycle_status": "Статус",
    "applies_target_class": "Класс цели",
    "applies_target_kinds": "Виды цели",
    "applies_implementation_scope": "Scope",
    "applies_dams_model_level": "Уровень DAMS",
    "applies_implementation_profile": "Профиль",
    "check_count": "Число проверок",
    "check_ids": "Check IDs",
    "check_kinds": "Виды проверок",
    "severities": "Severity",
    "diagnostic_codes": "Diagnostic codes",
}

_COLUMN_WIDTHS: dict[str, float] = {
    "catalog_id": 28,
    "catalog_name": 22,
    "source_path": 36,
    "element_id": 28,
    "code": 12,
    "name": 28,
    "title": 32,
    "description": 36,
    "statement": 64,
    "requirement_level": 16,
    "requirement_section": 10,
    "lifecycle_status": 12,
    "applies_target_class": 18,
    "applies_target_kinds": 16,
    "applies_implementation_scope": 12,
    "applies_dams_model_level": 14,
    "applies_implementation_profile": 14,
    "check_count": 10,
    "check_ids": 36,
    "check_kinds": 28,
    "severities": 18,
    "diagnostic_codes": 36,
}

_WRAP_COLUMNS = frozenset(
    {
        "title",
        "description",
        "statement",
        "source_path",
        "check_ids",
        "check_kinds",
        "diagnostic_codes",
    }
)

_JOIN = "; "

# Keep old name as alias for callers/tests that imported CSV_COLUMNS.
CSV_COLUMNS = EXPORT_COLUMNS


def _s(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return _JOIN.join(_s(v) for v in value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def _rel(path: Path, root: Path | None) -> str:
    if root is not None:
        try:
            return path.resolve().relative_to(root.resolve()).as_posix()
        except ValueError:
            pass
    return path.resolve().as_posix()


def requirement_row(
    req: dict[str, Any],
    *,
    catalog_id: str,
    catalog_name: str,
    source_path: str,
) -> dict[str, str]:
    applies = req.get("applies_to") if isinstance(req.get("applies_to"), dict) else {}
    checks = [c for c in (req.get("formal_checks") or []) if isinstance(c, dict)]
    row = {
        "catalog_id": catalog_id,
        "catalog_name": catalog_name,
        "source_path": source_path,
        "element_id": _s(req.get("element_id")),
        "code": _s(req.get("code")),
        "name": _s(req.get("name")),
        "title": _s(req.get("title")),
        "description": _s(req.get("description")),
        "statement": _s(req.get("statement")),
        "requirement_level": _s(req.get("requirement_level")),
        "requirement_section": _s(req.get("requirement_section")),
        "lifecycle_status": _s(req.get("lifecycle_status")),
        "applies_target_class": _s(applies.get("applies_target_class")),
        "applies_target_kinds": _s(applies.get("applies_target_kinds") or []),
        "applies_implementation_scope": _s(applies.get("applies_implementation_scope")),
        "applies_dams_model_level": _s(applies.get("applies_dams_model_level")),
        "applies_implementation_profile": _s(
            applies.get("applies_implementation_profile")
        ),
        "check_count": str(len(checks)),
        "check_ids": _s([c.get("check_id") for c in checks]),
        "check_kinds": _s([c.get("kind") for c in checks]),
        "severities": _s([c.get("severity") for c in checks]),
        "diagnostic_codes": _s([c.get("diagnostic_code") for c in checks]),
    }
    return {k: row[k] for k in EXPORT_COLUMNS}


def iter_catalog_rows(path: Path, *, source_path: str) -> list[dict[str, str]]:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError(f"{path} is not a YAML mapping")
    reqs = raw.get("requirements") or []
    if not isinstance(reqs, list):
        raise ValueError(f"{path} requirements is not a list")
    catalog_id = _s(raw.get("catalog_id"))
    catalog_name = _s(raw.get("name"))
    rows = [
        requirement_row(
            req,
            catalog_id=catalog_id,
            catalog_name=catalog_name,
            source_path=source_path,
        )
        for req in reqs
        if isinstance(req, dict)
    ]
    return sorted(
        rows,
        key=lambda r: (
            r["requirement_level"],
            r["requirement_section"],
            r["code"],
            r["element_id"],
        ),
    )


def default_catalog_paths(root: Path) -> list[Path]:
    folder = root / "model-assets/specifications/moex-dams/0.1/requirements"
    return sorted(p for p in folder.glob("*.yaml") if p.is_file())


def collect_requirement_rows(
    paths: list[Path],
    *,
    root: Path | None = None,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in paths:
        rows.extend(iter_catalog_rows(path, source_path=_rel(path, root)))
    return sorted(
        rows,
        key=lambda r: (
            r["requirement_level"],
            r["requirement_section"],
            r["code"],
            r["element_id"],
        ),
    )


def write_requirements_xlsx(rows: list[dict[str, str]], out: Path) -> int:
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "requirements"

    header_font = Font(bold=True)
    header_align = Alignment(vertical="center", wrap_text=True)
    wrap_align = Alignment(vertical="top", wrap_text=True)
    top_align = Alignment(vertical="top", wrap_text=False)

    for col_idx, key in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=HEADER_LABELS.get(key, key))
        cell.font = header_font
        cell.alignment = header_align

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            cell.alignment = wrap_align if key in _WRAP_COLUMNS else top_align

    for col_idx, key in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = _COLUMN_WIDTHS.get(
            key, 16
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for row_idx in range(2, len(rows) + 2):
        ws.row_dimensions[row_idx].height = 45

    wb.save(out)
    return len(rows)


# Back-compat alias (CSV path removed).
write_requirements_csv = write_requirements_xlsx
