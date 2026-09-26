"""Add Russian label/definition columns to fibo_definitions_client.xlsx."""

from __future__ import annotations

import json
import re
import time
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from deep_translator import MyMemoryTranslator
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "output" / "fibo_definitions_client.xlsx"
OUT = ROOT / "output" / "fibo_definitions_client.xlsx"
CACHE = ROOT / "output" / "_client_ru_cache.json"
MAX_W = 80

# High-quality overrides for core glossary terms (label / short phrase)
LABEL_RU: dict[str, str] = {
    "account holder": "держатель счёта",
    "all borrowers' monthly income": "совокупный ежемесячный доход всех заёмщиков",
    "beneficial owner": "бенефициарный владелец",
    "beneficial ownership": "бенефициарное владение",
    "borrower": "заёмщик",
    "borrower assessment": "оценка заёмщика",
    "borrower data protection requirement": "требование о защите данных заёмщика",
    "borrower disclosure requirement": "требование о раскрытии информации заёмщиком",
    "borrower identification scheme": "схема идентификации заёмщика",
    "borrower identifier": "идентификатор заёмщика",
    "borrower monthly income": "ежемесячный доход заёмщика",
    "borrower right": "право заёмщика",
    "cardholder": "держатель карты",
    "central counterparty clearing house": "центральный контрагент (клиринговая организация)",
    "client": "клиент",
    "client identifier": "идентификатор клиента",
    "counterparty": "контрагент",
    "customer": "клиент (customer)",
    "customer account": "клиентский счёт",
    "customer account holder": "держатель клиентского счёта",
    "customer identifier": "идентификатор клиента (customer)",
    "futures trading account holder": "держатель счёта для торговли фьючерсами",
    "loan-specific customer account": "клиентский счёт по кредиту",
    "policyholder": "страхователь",
    "receiving counterparty": "принимающий контрагент",
    "securities transaction counterparty": "контрагент по сделке с ценными бумагами",
    "transaction counterparty": "контрагент по сделке",
    "has beneficial owner": "имеет бенефициарного владельца",
    "has borrower": "имеет заёмщика",
    "has counterparty": "имеет контрагента",
    "has primary account holder": "имеет основного держателя счёта",
    "has secondary account holder": "имеет дополнительного держателя счёта",
    "is beneficial owner of": "является бенефициарным владельцем",
    "investor": "инвестор",
    "primary investor": "первичный инвестор",
    "obligor": "должник (obligor)",
    "relationship manager": "менеджер по работе с клиентами",
    "service agreement": "соглашение об обслуживании",
    "account-specific service agreement": "соглашение об обслуживании конкретного счёта",
    "depositor": "вкладчик",
}

# Split long texts for MyMemory (limit ~500 chars)
_CHUNK = 450


def _translate_mt(text: str, translator: MyMemoryTranslator) -> str:
    text = text.strip()
    if not text:
        return ""
    if len(text) <= _CHUNK:
        return translator.translate(text)
    # split on sentence boundaries roughly
    parts: list[str] = []
    buf = ""
    for sent in re.split(r"(?<=[.!?])\s+", text):
        if len(buf) + len(sent) + 1 <= _CHUNK:
            buf = f"{buf} {sent}".strip()
        else:
            if buf:
                parts.append(translator.translate(buf))
                time.sleep(0.35)
            buf = sent
    if buf:
        parts.append(translator.translate(buf))
        time.sleep(0.35)
    return " ".join(parts)


def translate_label(label: str | None, cache: dict, translator: MyMemoryTranslator) -> str | None:
    if not label or (isinstance(label, float) and pd.isna(label)):
        return None
    key = str(label).strip()
    low = key.lower()
    if low in LABEL_RU:
        return LABEL_RU[low]
    # multi-label "a | b"
    if " | " in key:
        return " | ".join(
            translate_label(p, cache, translator) or p for p in key.split(" | ")
        )
    ck = f"label::{key}"
    if ck in cache:
        return cache[ck]
    try:
        ru = _translate_mt(key, translator)
        time.sleep(0.3)
    except Exception as exc:  # noqa: BLE001
        print(f"LABEL FAIL [{key}]: {exc}")
        ru = None
    cache[ck] = ru
    return ru


def translate_definition(
    definition: str | None, cache: dict, translator: MyMemoryTranslator
) -> str | None:
    if not definition or (isinstance(definition, float) and pd.isna(definition)):
        return None
    key = str(definition).strip()
    # keep multiple definitions separated
    if "\n---\n" in key:
        parts = key.split("\n---\n")
        ru_parts = [translate_definition(p, cache, translator) or p for p in parts]
        return "\n---\n".join(ru_parts)
    ck = f"def::{key}"
    if ck in cache:
        return cache[ck]
    try:
        ru = _translate_mt(key, translator)
        time.sleep(0.35)
    except Exception as exc:  # noqa: BLE001
        print(f"DEF FAIL [{key[:60]}...]: {exc}")
        ru = None
    cache[ck] = ru
    return ru


def style(ws) -> None:
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = ws.cell(1, col_idx).value
        max_len = len(str(header or ""))
        for row_idx in range(2, min(ws.max_row + 1, 250)):
            val = ws.cell(row_idx, col_idx).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), MAX_W))
        ws.column_dimensions[letter].width = min(max_len + 2, MAX_W)
    for col_idx in range(1, ws.max_column + 1):
        if ws.cell(1, col_idx).value in {
            "definition",
            "определение по русски",
            "термин по русски",
        }:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, col_idx).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )


def enrich(df: pd.DataFrame, cache: dict, translator: MyMemoryTranslator) -> pd.DataFrame:
    out = df.copy()
    labels_ru = []
    defs_ru = []
    for i, row in out.iterrows():
        print(f"  [{i+1}/{len(out)}] {row.get('local_name')}")
        labels_ru.append(translate_label(row.get("label"), cache, translator))
        defs_ru.append(translate_definition(row.get("definition"), cache, translator))
        # persist cache often
        if (i + 1) % 10 == 0:
            CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    # insert Russian columns right after English label/definition
    cols = list(out.columns)
    out["термин по русски"] = labels_ru
    out["определение по русски"] = defs_ru
    # reorder: after label put термин, after definition put определение
    new_cols: list[str] = []
    for c in cols:
        new_cols.append(c)
        if c == "label":
            new_cols.append("термин по русски")
        if c == "definition":
            new_cols.append("определение по русски")
    # dedupe if already present
    seen = set()
    ordered = []
    for c in new_cols:
        if c not in seen and c in out.columns:
            seen.add(c)
            ordered.append(c)
    for c in out.columns:
        if c not in seen:
            ordered.append(c)
    return out[ordered]


def main() -> None:
    cache: dict = {}
    if CACHE.exists():
        cache = json.loads(CACHE.read_text(encoding="utf-8"))
        print(f"Loaded cache: {len(cache)} entries")

    translator = MyMemoryTranslator(source="en-GB", target="ru-RU")

    sheets = {
        "Core client": pd.read_excel(SRC, sheet_name="Core client"),
        "Glossary": pd.read_excel(SRC, sheet_name="Glossary"),
        "All related": pd.read_excel(SRC, sheet_name="All related"),
    }
    readme = pd.read_excel(SRC, sheet_name="README")

    enriched = {}
    # Translate unique labels/defs once via All related (superset), then map
    print("Translating All related (superset)...")
    enriched["All related"] = enrich(sheets["All related"], cache, translator)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    # Build lookup from all related
    by_iri = {}
    for _, row in enriched["All related"].iterrows():
        by_iri[row["iri"]] = (
            row["термин по русски"],
            row["определение по русски"],
        )

    def apply_lookup(df: pd.DataFrame) -> pd.DataFrame:
        out = df.copy()
        labels_ru, defs_ru = [], []
        for _, row in out.iterrows():
            pair = by_iri.get(row["iri"])
            if pair:
                labels_ru.append(pair[0])
                defs_ru.append(pair[1])
            else:
                labels_ru.append(translate_label(row.get("label"), cache, translator))
                defs_ru.append(
                    translate_definition(row.get("definition"), cache, translator)
                )
        cols = list(out.columns)
        out["термин по русски"] = labels_ru
        out["определение по русски"] = defs_ru
        new_cols: list[str] = []
        for c in cols:
            new_cols.append(c)
            if c == "label":
                new_cols.append("термин по русски")
            if c == "definition":
                new_cols.append("определение по русски")
        seen = set()
        ordered = []
        for c in new_cols:
            if c not in seen and c in out.columns:
                seen.add(c)
                ordered.append(c)
        for c in out.columns:
            if c not in seen:
                ordered.append(c)
        return out[ordered]

    print("Applying to Core client / Glossary...")
    enriched["Core client"] = apply_lookup(sheets["Core client"])
    enriched["Glossary"] = apply_lookup(sheets["Glossary"])

    # Update README
    extra = pd.DataFrame(
        [
            {
                "property": "ru_columns_added_at_utc",
                "value": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            },
            {
                "property": "ru_columns",
                "value": "термин по русски; определение по русски",
            },
            {
                "property": "ru_translation_note",
                "value": (
                    "Machine translation (MyMemory) with manual overrides for core "
                    "client/customer terms; review recommended before publication."
                ),
            },
        ]
    )
    readme2 = pd.concat([readme, extra], ignore_index=True)

    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        readme2.to_excel(writer, sheet_name="README", index=False)
        enriched["Core client"].to_excel(writer, sheet_name="Core client", index=False)
        enriched["Glossary"].to_excel(writer, sheet_name="Glossary", index=False)
        enriched["All related"].to_excel(writer, sheet_name="All related", index=False)
        for name in ["README", "Core client", "Glossary", "All related"]:
            style(writer.sheets[name])

    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
