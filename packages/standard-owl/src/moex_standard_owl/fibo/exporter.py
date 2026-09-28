"""Write FIBO definition marts: CSV, Excel, JSON, manifest."""

from __future__ import annotations

import json
import logging
import platform
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError, version as pkg_version
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from moex_standard_owl.fibo.constants import UPSTREAM_URL
from moex_standard_owl.models import (
    CORE_COLUMNS,
    PARSE_ERROR_COLUMNS,
    TECHNICAL_COLUMNS,
    AggregatedEntity,
    ExportManifest,
    ParseError,
)

logger = logging.getLogger(__name__)

MAX_COL_WIDTH = 80


def _lib_version(name: str) -> str:
    try:
        return pkg_version(name)
    except PackageNotFoundError:
        return "unknown"


def entities_to_dataframe(
    entities: list[AggregatedEntity],
    *,
    technical: bool = False,
) -> pd.DataFrame:
    columns = list(TECHNICAL_COLUMNS if technical else CORE_COLUMNS)
    if not entities:
        return pd.DataFrame(columns=columns)
    rows = [
        e.to_technical_dict() if technical else e.to_core_dict() for e in entities
    ]
    df = pd.DataFrame(rows)
    return df.reindex(columns=columns)


def parse_errors_to_dataframe(errors: list[ParseError]) -> pd.DataFrame:
    if not errors:
        return pd.DataFrame(columns=list(PARSE_ERROR_COLUMNS))
    rows = [
        {
            "module_path": e.module_path,
            "exception_type": e.exception_type,
            "exception_message": e.exception_message,
        }
        for e in errors
    ]
    return pd.DataFrame(rows).reindex(columns=list(PARSE_ERROR_COLUMNS))


def write_csv(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8")
    logger.info("Wrote %s (%d rows)", path, len(df))


def write_json_entities(entities: list[AggregatedEntity], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = [e.to_technical_dict() for e in entities]
    with path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    logger.info("Wrote %s (%d entities)", path, len(payload))


def write_manifest(manifest: ExportManifest, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(manifest.to_dict(), fh, ensure_ascii=False, indent=2)
    logger.info("Wrote %s", path)


def _style_worksheet(ws, *, wrap_definition: bool = True) -> None:
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    # Column widths
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = ws.cell(1, col_idx).value
        max_len = len(str(header)) if header else 10
        for row_idx in range(2, min(ws.max_row + 1, 200)):  # sample first rows
            val = ws.cell(row_idx, col_idx).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), MAX_COL_WIDTH))
        ws.column_dimensions[letter].width = min(max_len + 2, MAX_COL_WIDTH)

    if wrap_definition:
        def_col = None
        for col_idx in range(1, ws.max_column + 1):
            if ws.cell(1, col_idx).value == "definition":
                def_col = col_idx
                break
        if def_col is not None:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, def_col).alignment = Alignment(wrap_text=True, vertical="top")


def _write_df_sheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    # Excel sheet name max 31 chars
    name = sheet_name[:31]
    df.to_excel(writer, sheet_name=name, index=False)
    _style_worksheet(writer.sheets[name])


def build_readme_rows(
    *,
    generated_at_utc: str,
    source_path: str,
    source_release: str,
    domains: list[str],
    types: list[str],
    processed_files: int,
    failed_files: int,
    sheet_counts: dict[str, int],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = [
        {"property": "generated_at_utc", "value": generated_at_utc},
        {"property": "source_path", "value": source_path},
        {"property": "source_release", "value": source_release},
        {"property": "included_domains", "value": ", ".join(domains)},
        {"property": "included_types", "value": ", ".join(types)},
        {"property": "processed_files", "value": processed_files},
        {"property": "failed_files", "value": failed_files},
        {"property": "python_version", "value": platform.python_version()},
        {"property": "rdflib_version", "value": _lib_version("rdflib")},
        {"property": "pandas_version", "value": _lib_version("pandas")},
        {"property": "openpyxl_version", "value": _lib_version("openpyxl")},
        {"property": "upstream", "value": UPSTREAM_URL},
    ]
    for sheet, count in sheet_counts.items():
        rows.append({"property": f"sheet.{sheet}.rows", "value": count})
    return pd.DataFrame(rows)


def write_excel(
    path: Path,
    *,
    readme_df: pd.DataFrame,
    glossary_df: pd.DataFrame,
    all_df: pd.DataFrame,
    active_df: pd.DataFrame,
    classes_df: pd.DataFrame,
    properties_df: pd.DataFrame,
    individuals_df: pd.DataFrame,
    deprecated_df: pd.DataFrame,
    no_definition_df: pd.DataFrame,
    modules_df: pd.DataFrame,
    parse_errors_df: pd.DataFrame,
    technical_df: pd.DataFrame,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_df_sheet(writer, "README", readme_df)
        _write_df_sheet(writer, "Glossary", glossary_df)
        _write_df_sheet(writer, "All entities", all_df)
        _write_df_sheet(writer, "Active entities", active_df)
        _write_df_sheet(writer, "Classes", classes_df)
        _write_df_sheet(writer, "Properties", properties_df)
        _write_df_sheet(writer, "Individuals", individuals_df)
        _write_df_sheet(writer, "Deprecated", deprecated_df)
        _write_df_sheet(writer, "No definition", no_definition_df)
        _write_df_sheet(writer, "Modules summary", modules_df)
        _write_df_sheet(writer, "Parse errors", parse_errors_df)
        _write_df_sheet(writer, "Technical details", technical_df)
    logger.info("Wrote %s", path)


def modules_summary(entities: list[AggregatedEntity]) -> pd.DataFrame:
    """Count entities per module_path segment (split aggregated paths)."""
    counts: dict[str, int] = {}
    for ent in entities:
        if not ent.module_path:
            continue
        for part in ent.module_path.split(" | "):
            part = part.strip()
            if part:
                counts[part] = counts.get(part, 0) + 1
    rows = [
        {"module_path": path, "entity_count": count}
        for path, count in sorted(counts.items())
    ]
    if not rows:
        return pd.DataFrame(columns=["module_path", "entity_count"])
    return pd.DataFrame(rows)


PROPERTY_TYPES = frozenset(
    {"owl:ObjectProperty", "owl:DatatypeProperty", "owl:AnnotationProperty"}
)


def split_views(
    main: list[AggregatedEntity],
    all_aggregated: list[AggregatedEntity],
    *,
    include_deprecated_in_main: bool,
) -> dict[str, list[AggregatedEntity]]:
    """Build named entity lists for CSV/Excel views."""
    # all / active for primary marts use `main` (already filtered)
    active = [e for e in main if not e.is_deprecated]
    deprecated = [e for e in all_aggregated if e.is_deprecated]
    no_definition = [e for e in main if e.definition is None]
    classes = [e for e in main if e.entity_type == "owl:Class"]
    properties = [e for e in main if e.entity_type in PROPERTY_TYPES]
    # Individuals sheet: always from full aggregated set of individuals
    individuals = [
        e for e in all_aggregated if e.entity_type == "owl:NamedIndividual"
    ]
    # Glossary mart: active classes that have a formal definition
    glossary = [
        e
        for e in main
        if e.entity_type == "owl:Class"
        and not e.is_deprecated
        and e.definition is not None
    ]
    return {
        "all": main,
        "active": active,
        "deprecated": deprecated,
        "no_definition": no_definition,
        "classes": classes,
        "properties": properties,
        "individuals": individuals,
        "glossary": glossary,
    }


GLOSSARY_COLUMNS: tuple[str, ...] = (
    "local_name",
    "label",
    "definition",
    "iri",
    "module_path",
    "source_domain",
    "source_release",
    "source_file_sha256",
)


def entities_to_glossary_dataframe(entities: list[AggregatedEntity]) -> pd.DataFrame:
    """Compact class glossary suitable for corporate-term mapping."""
    if not entities:
        return pd.DataFrame(columns=list(GLOSSARY_COLUMNS))
    rows = []
    for e in entities:
        rows.append(
            {
                "local_name": e.local_name,
                "label": e.label,
                "definition": e.definition,
                "iri": e.iri,
                "module_path": e.module_path,
                "source_domain": e.source_domain,
                "source_release": e.source_release,
                "source_file_sha256": e.source_file_sha256,
            }
        )
    return pd.DataFrame(rows).reindex(columns=list(GLOSSARY_COLUMNS))


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def run_export(
    *,
    output_dir: Path,
    source_path: Path,
    source_release: str,
    domains: list[str],
    type_codes: list[str],
    include_examples: bool,
    include_individuals: bool,
    include_deprecated: bool,
    formats: list[str],
    main_entities: list[AggregatedEntity],
    all_aggregated: list[AggregatedEntity],
    parse_errors: list[ParseError],
    processed_files: int,
    successful_files: int,
) -> ExportManifest:
    views = split_views(
        main_entities,
        all_aggregated,
        include_deprecated_in_main=include_deprecated,
    )

    all_df = entities_to_dataframe(views["all"])
    active_df = entities_to_dataframe(views["active"])
    deprecated_df = entities_to_dataframe(views["deprecated"])
    no_def_df = entities_to_dataframe(views["no_definition"])
    classes_df = entities_to_dataframe(views["classes"])
    properties_df = entities_to_dataframe(views["properties"])
    individuals_df = entities_to_dataframe(views["individuals"])
    glossary_df = entities_to_glossary_dataframe(views["glossary"])
    technical_df = entities_to_dataframe(all_aggregated, technical=True)
    errors_df = parse_errors_to_dataframe(parse_errors)
    modules_df = modules_summary(all_aggregated)

    output_files: list[str] = []
    generated_at = utc_now_iso()

    if "csv" in formats:
        mapping = {
            "fibo_glossary.csv": glossary_df,
            "fibo_definitions_all.csv": all_df,
            "fibo_definitions_active.csv": active_df,
            "fibo_definitions_deprecated.csv": deprecated_df,
            "fibo_definitions_no_definition.csv": no_def_df,
            "fibo_parse_errors.csv": errors_df,
        }
        for name, df in mapping.items():
            write_csv(df, output_dir / name)
            output_files.append(name)

    if "xlsx" in formats:
        sheet_counts = {
            "Glossary": len(glossary_df),
            "All entities": len(all_df),
            "Active entities": len(active_df),
            "Classes": len(classes_df),
            "Properties": len(properties_df),
            "Individuals": len(individuals_df),
            "Deprecated": len(deprecated_df),
            "No definition": len(no_def_df),
            "Modules summary": len(modules_df),
            "Parse errors": len(errors_df),
            "Technical details": len(technical_df),
        }
        readme_df = build_readme_rows(
            generated_at_utc=generated_at,
            source_path=str(source_path),
            source_release=source_release,
            domains=domains,
            types=type_codes,
            processed_files=processed_files,
            failed_files=len(parse_errors),
            sheet_counts=sheet_counts,
        )
        xlsx_name = "fibo_definitions.xlsx"
        write_excel(
            output_dir / xlsx_name,
            readme_df=readme_df,
            glossary_df=glossary_df,
            all_df=all_df,
            active_df=active_df,
            classes_df=classes_df,
            properties_df=properties_df,
            individuals_df=individuals_df,
            deprecated_df=deprecated_df,
            no_definition_df=no_def_df,
            modules_df=modules_df,
            parse_errors_df=errors_df,
            technical_df=technical_df,
        )
        output_files.append(xlsx_name)

    if "json" in formats:
        json_name = "fibo_definitions_all.json"
        write_json_entities(views["all"], output_dir / json_name)
        output_files.append(json_name)

    # Always write manifest
    manifest_name = "fibo_export_manifest.json"
    entities_deprecated = sum(1 for e in all_aggregated if e.is_deprecated)
    entities_active = sum(1 for e in main_entities if not e.is_deprecated)
    entities_no_def = sum(1 for e in main_entities if e.definition is None)

    manifest = ExportManifest(
        generated_at_utc=generated_at,
        source_path=str(source_path.resolve()),
        source_release=source_release,
        included_domains=list(domains),
        included_types=list(type_codes),
        include_examples=include_examples,
        include_individuals=include_individuals,
        include_deprecated=include_deprecated,
        processed_files=processed_files,
        successful_files=successful_files,
        failed_files=len(parse_errors),
        entities_total=len(main_entities),
        entities_active=entities_active,
        entities_deprecated=entities_deprecated,
        entities_without_definition=entities_no_def,
        output_files=output_files + [manifest_name],
    )
    write_manifest(manifest, output_dir / manifest_name)
    return manifest
