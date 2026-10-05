"""CLI adapter tests — handlers, not HTML."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from moex_model_cli.__main__ import main
from moex_model_cli.bootstrap import (
    DEFAULT_IMPLEMENTATION_ENVELOPE,
    DEFAULT_SPECIFICATION_ENVELOPE,
    SlicePaths,
    find_repo_root,
    resolve_body_from_implementation_envelope,
    resolve_schema_from_specification_envelope,
)


def test_find_repo_root(repo_root: Path) -> None:
    assert find_repo_root(repo_root) == repo_root
    nested = repo_root / "apps" / "cli"
    assert find_repo_root(nested) == repo_root


def test_envelopes_resolve_schema_and_body(repo_root: Path) -> None:
    spec_env = repo_root / DEFAULT_SPECIFICATION_ENVELOPE
    impl_env = repo_root / DEFAULT_IMPLEMENTATION_ENVELOPE
    assert spec_env.is_file()
    assert impl_env.is_file()
    schema = resolve_schema_from_specification_envelope(spec_env)
    body = resolve_body_from_implementation_envelope(impl_env)
    assert schema.name == "moex-dams.yaml"
    assert body.name == "mdm-solution-model.yaml"
    paths = SlicePaths.resolve(root=repo_root)
    assert paths.schema == schema
    assert paths.implementation == body


def test_validate_mdm_solution(repo_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = main(["validate", "--root", str(repo_root)])
    captured = capsys.readouterr()
    assert code == 0
    assert "overall=conformant_with_warnings" in captured.out
    assert "dams:model/mdm/0.1.0" in captured.out


def test_validate_json(repo_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = main(["validate", "--root", str(repo_root), "--json"])
    captured = capsys.readouterr()
    assert code == 0
    payload = json.loads(captured.out)
    assert payload["overall_result"] == "conformant_with_warnings"
    assert payload["assessed_specification"]["specification_id"] == "moex:spec:dams"


def test_publish_writes_json(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    out = tmp_path / "slice.json"
    code = main(
        [
            "publish",
            "--root",
            str(repo_root),
            "--out",
            str(out),
            "--implementation-id",
            "moex:implementation:mdm:0.1.0",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0
    assert out.is_file()
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["summary"]["is_conformant"] is True
    assert data["summary"]["package_id"] == "dams:model/mdm/0.1.0"
    assert "wrote" in captured.out


def test_lint_and_compile(repo_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["lint", "--root", str(repo_root)]) == 0
    assert "lint diagnostics=" in capsys.readouterr().out
    assert main(["compile", "--root", str(repo_root)]) == 0
    assert "contracts OK" in capsys.readouterr().out


def test_compile_artifacts(repo_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["compile", "--root", str(repo_root), "--artifacts"]) == 0
    out = capsys.readouterr().out
    assert "contracts OK" in out
    assert "artifacts OK" in out
    assert (repo_root / "generated" / "artifacts" / "moex-dams" / "0.1" / "moex-dams.dbml").is_file()
    assert (repo_root / "generated" / "manifests" / "moex-dams-owl.json").is_file()


def test_diagram_projects_logical(
    repo_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out.dbml"
    assert main(
        [
            "diagram",
            "--root",
            str(repo_root),
            "--profile",
            "logical",
            "--out",
            str(out),
        ]
    ) == 0
    text = out.read_text(encoding="utf-8")
    assert "Table ENTERPRISE" in text
    assert "ENTERPRISE_ID" in text
    assert (tmp_path / "out.dbml.manifest.json").is_file()
    assert "digest=sha256:" in capsys.readouterr().out


def test_diagram_projects_conceptual_dbml(
    repo_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    impl = (
        repo_root
        / "model-assets"
        / "implementations"
        / "enterprise"
        / "moex-enterprise-conceptual-model"
        / "0.1"
        / "enterprise-conceptual-model.yaml"
    )
    out = tmp_path / "conceptual.dbml"
    assert main(
        [
            "diagram",
            "--root",
            str(repo_root),
            "--implementation",
            str(impl),
            "--profile",
            "conceptual",
            "--format",
            "dbml",
            "--out",
            str(out),
        ]
    ) == 0
    text = out.read_text(encoding="utf-8")
    assert "Table Organization" in text
    assert "headercolor: #F4B400" in text
    assert "concept string" in text
    assert "digest=sha256:" in capsys.readouterr().out


def test_diagram_projects_mermaid_logical(
    repo_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "logical.erd.md"
    assert main(
        [
            "diagram",
            "--root",
            str(repo_root),
            "--profile",
            "logical",
            "--format",
            "mermaid",
            "--no-svg",
            "--out",
            str(out),
        ]
    ) == 0
    text = out.read_text(encoding="utf-8")
    assert "```mermaid" in text
    assert "erDiagram" in text
    assert "ENTERPRISE" in text
    assert "digest=sha256:" in capsys.readouterr().out
    scene = out.parent / "logical.scene.json"
    layout = out.parent / "logical.layout.json"
    assert scene.is_file()
    assert layout.is_file()


def test_import_er_dictionary_fixture(
    repo_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = (
        repo_root / "packages" / "standard-linkml" / "tests" / "fixtures" / "er-dictionary"
    )
    profile = workbook / "profile.yaml"
    out = tmp_path / "ingest-out"
    code = main(
        [
            "import",
            "--root",
            str(repo_root),
            "--workbook",
            str(workbook),
            "--profile",
            str(profile),
            "--out",
            str(out),
            "--skip-validate",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0, captured.out + captured.err
    assert (out / "pilot_solution_model.package.yaml").is_file()
    assert "import OK" in captured.out


def test_map_sssom_and_extract(
    repo_root: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    sssom = (
        repo_root
        / "model-assets"
        / "transformations"
        / "mappings"
        / "dams-fibo.sssom.yaml"
    )
    assert main(["map", "--root", str(repo_root), "--sssom", str(sssom)]) == 0
    out = capsys.readouterr().out
    assert "moex:mappings:dams-fibo" in out
    assert "bindings=1" in out

    assert main(["map", "--root", str(repo_root)]) == 2
    assert "exactly one of" in capsys.readouterr().out

    fixture = (
        repo_root
        / "packages"
        / "semantic-mappings"
        / "tests"
        / "fixtures"
        / "mapped_schema.yaml"
    )
    assert (
        main(
            [
                "map",
                "--root",
                str(repo_root),
                "--extract-schema",
                str(fixture),
            ]
        )
        == 0
    )
    assert "bindings=" in capsys.readouterr().out


def test_map_transform_preview(
    repo_root: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    pytest.importorskip("linkml_map")
    transform = repo_root / "model-assets" / "transformations" / "person-rename-code.yaml"
    sample = (
        repo_root
        / "packages"
        / "linkml-tooling"
        / "tests"
        / "fixtures"
        / "person_sample.json"
    )
    code = main(
        [
            "map",
            "--root",
            str(repo_root),
            "--transform",
            str(transform),
            "--preview",
            str(sample),
        ]
    )
    out = capsys.readouterr().out
    assert code == 0, out
    assert "person-rename-code" in out or "transform preview" in out


def test_import_json_schema_draft(
    repo_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    pytest.importorskip("schema_automator")
    source = (
        repo_root
        / "packages"
        / "linkml-tooling"
        / "tests"
        / "fixtures"
        / "mini.schema.json"
    )
    out = tmp_path / "imports"
    code = main(
        [
            "import",
            "--root",
            str(repo_root),
            "--source-type",
            "json_schema",
            "--source",
            str(source),
            "--out",
            str(out),
            "--name",
            "MiniPerson",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0, captured.out + captured.err
    assert "generated-draft" in captured.out
    assert any((out / d / "job.json").is_file() for d in out.iterdir() if d.is_dir())


def test_missing_schema_fails_fast(repo_root: Path) -> None:
    missing = repo_root / "no-such-schema.yaml"
    paths = SlicePaths.resolve(root=repo_root, schema=missing)
    assert not paths.schema.is_file()
    with pytest.raises(FileNotFoundError):
        main(["validate", "--root", str(repo_root), "--schema", str(missing)])


def test_semantic_diff_identical_exit_0(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    mdm = (
        repo_root
        / "model-assets"
        / "implementations"
        / "solutions"
        / "mdm"
        / "mdm-solution-model.yaml"
    )
    left = tmp_path / "left.yaml"
    right = tmp_path / "right.yaml"
    text = mdm.read_text(encoding="utf-8")
    left.write_text(text, encoding="utf-8")
    right.write_text(text, encoding="utf-8")
    code = main(
        [
            "semantic-diff",
            "--root",
            str(repo_root),
            "--left",
            str(left),
            "--right",
            str(right),
            "--json",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0
    payload = json.loads(captured.out)
    assert payload["changes"] == []


def test_semantic_diff_breaking_exit_1(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    import yaml

    mdm = (
        repo_root
        / "model-assets"
        / "implementations"
        / "solutions"
        / "mdm"
        / "mdm-solution-model.yaml"
    )
    left = tmp_path / "left.yaml"
    right = tmp_path / "right.yaml"
    left.write_text(mdm.read_text(encoding="utf-8"), encoding="utf-8")
    data = yaml.safe_load(mdm.read_text(encoding="utf-8"))
    for entity in data["logical_entities"]:
        if entity.get("element_id") == "dams:logical/mdm/ENTERPRISE":
            entity["attributes"] = [
                a
                for a in entity["attributes"]
                if a.get("element_id") != "dams:logical/mdm/ENTERPRISE/SHORT_NAME"
            ]
    right.write_text(
        yaml.safe_dump(data, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    code = main(
        [
            "semantic-diff",
            "--root",
            str(repo_root),
            "--left",
            str(left),
            "--right",
            str(right),
            "--json",
        ]
    )
    captured = capsys.readouterr()
    assert code == 1
    payload = json.loads(captured.out)
    categories = {c["category"] for c in payload["changes"]}
    assert "breaking" in categories


def test_import_solution_fixture(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    xlsx = (
        repo_root
        / "packages"
        / "standard-linkml"
        / "tests"
        / "fixtures"
        / "solution-xlsx"
        / "fixture.xlsx"
    )
    profile = (
        repo_root
        / "packages"
        / "standard-linkml"
        / "tests"
        / "fixtures"
        / "solution-xlsx"
        / "profile.yaml"
    )
    out = tmp_path / "crm"
    code = main(
        [
            "import-solution",
            "--root",
            str(repo_root),
            "--xlsx",
            str(xlsx),
            "--profile",
            str(profile),
            "--system",
            "CRM",
            "--out",
            str(out),
            "--report-dir",
            str(tmp_path / "report"),
            "--force",
            "--skip-assess",
            "--skip-export",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0, captured.out
    assert (out / "crm-solution-model.yaml").is_file()
    assert (out / "implementation.yaml").is_file()
    assert (out / "publish.yaml").is_file()
    assert "import-solution" in captured.out


def test_export_requirements_xlsx(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    from openpyxl import load_workbook

    out = tmp_path / "requirements.xlsx"
    code = main(
        ["export-requirements", "--root", str(repo_root), "--out", str(out)]
    )
    captured = capsys.readouterr()
    assert code == 0, captured.out
    assert out.is_file()
    assert "rows=34" in captured.out
    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "requirements"
    codes = {ws.cell(row=r, column=5).value for r in range(2, ws.max_row + 1)}
    assert "GEN-001" in codes
    assert "CM-GEN-001" in codes
    statements = {
        ws.cell(row=r, column=9).value or "" for r in range(2, ws.max_row + 1)
    }
    assert any("У пакета модели данных ИТ-решения" in s for s in statements)


def test_ontology_report_json(
    repo_root: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    code = main(["ontology-report", "--root", str(repo_root), "--json"])
    captured = capsys.readouterr()
    assert code == 0, captured.out
    payload = json.loads(captured.out)
    assert payload["schema_name"] == "moex_dams"
    assert payload["default_prefix"] == "dams"
    assert payload["issues"] == []


def test_ontology_report_uri_diff_breaking(
    repo_root: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    old = tmp_path / "old.yaml"
    new = tmp_path / "new.yaml"
    old.write_text(
        """
id: https://example.com/old
name: old
prefixes:
  ex: https://example.com/
default_prefix: ex
classes:
  Person:
    attributes:
      name:
        range: string
""",
        encoding="utf-8",
    )
    new.write_text(
        """
id: https://example.com/new
name: new
prefixes:
  ex: https://example.com/
default_prefix: ex
classes:
  Human:
    attributes:
      name:
        range: string
""",
        encoding="utf-8",
    )
    code = main(
        [
            "ontology-report",
            "--root",
            str(repo_root),
            "--from",
            str(old),
            "--to",
            str(new),
        ]
    )
    captured = capsys.readouterr()
    assert code == 1
    assert "MOEX-ONT-010" in captured.out
    assert "BREAKING" in captured.out
